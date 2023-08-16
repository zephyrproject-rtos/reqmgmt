import argparse
import os
import string
from enum import Enum
from pathlib import Path
from typing import List

import openpyxl
from strictdoc.backend.sdoc.document_reference import DocumentReference
from strictdoc.backend.sdoc.models.document import Document
from strictdoc.backend.sdoc.models.document_config import DocumentConfig
from strictdoc.backend.sdoc.models.document_grammar import DocumentGrammar, GrammarElement
from strictdoc.backend.sdoc.models.object_factory import SDocObjectFactory
from strictdoc.backend.sdoc.models.reference import Reference, ParentReqReference
from strictdoc.backend.sdoc.models.requirement import Requirement, RequirementField
from strictdoc.backend.sdoc.models.type_system import RequirementFieldName, GrammarElementFieldString, \
    GrammarReferenceType, GrammarElementFieldReference
from strictdoc.backend.sdoc.writer import SDWriter


def letter_index(letter: str):
    assert isinstance(letter, str)
    assert len(letter) == 1
    return string.ascii_uppercase.index(letter)


def main():
    # Step: Parse command-line argument (input file).
    parser = argparse.ArgumentParser()

    parser.add_argument("input_file", type=str, help="TODO")
    args = parser.parse_args()

    print(args.input_file)  # noqa: T201

    path_to_input_html = args.input_file
    if not os.path.isfile(path_to_input_html):
        print(f"not a file: {path_to_input_html}")  # noqa: T201
        exit(1)

    # Step: Parse Excel content into a list of Python dictionaries, one
    # dictionary per requirement.
    excel_workbook = openpyxl.load_workbook(path_to_input_html, read_only=True)
    worksheet_with_requirements = excel_workbook["Zephyr Requirements"]

    class FIELDS(str, Enum):
        UID = "UID"
        PARENT_UID = "PARENT_UID"
        COMPONENT = "COMPONENT"
        TYPE = "TYPE"
        TITLE = "TITLE"
        STATEMENT = "STATEMENT"
        USER_STORY = "USER_STORY"
        DISCUSSION_DATE = "DISCUSSION_DATE"
        REVIEW_COMMENT = "REVIEW_COMMENT"

    requirements_dicts = []
    for row in worksheet_with_requirements.iter_rows(
        min_row=2, max_col=20, max_row=122
    ):
        # print(row)

        requirement_dict = {
            FIELDS.UID.value: row[letter_index("C")].value,
            FIELDS.PARENT_UID.value: row[letter_index("D")].value,
            FIELDS.COMPONENT.value: row[letter_index("F")].value,
            FIELDS.TITLE.value: row[letter_index("S")].value,
            FIELDS.STATEMENT.value: row[letter_index("M")].value,
            FIELDS.TYPE.value: row[letter_index("G")].value,
            FIELDS.USER_STORY.value: row[letter_index("N")].value,
            FIELDS.DISCUSSION_DATE.value: row[letter_index("A")].value,
            FIELDS.REVIEW_COMMENT.value: row[letter_index("O")].value
        }
        requirements_dicts.append(requirement_dict)

    first_requirement_dict = requirements_dicts[0]
    assert first_requirement_dict[FIELDS.UID.value] == "ZEP-ARCH-001"
    assert first_requirement_dict[FIELDS.STATEMENT.value] == "Zephyr shall provide a framework for applications running on Zephyr to communicate with the hardware architectural service available from Zephyr for the system"

    # print(requirements_dicts)

    # Step: Convert Python dictionaries with requirements to StrictDoc requirements.
    document = Document(
        title="Zephyr Requirements",
        config=None,
        grammar=None,
        free_texts=[],
        section_contents=[]
    )
    document.config = DocumentConfig(
        parent=document, version=None, uid=None, classification=None, requirement_prefix=None, markup="Text", auto_levels=None, requirement_style=None, requirement_in_toc=None
    )

    fields = [
        GrammarElementFieldString(
            parent=None, title=RequirementFieldName.UID, required="False"
        ),
        GrammarElementFieldString(
            parent=None,
            title="TYPE",
            required="False",
        ),
        GrammarElementFieldString(
            parent=None,
            title=FIELDS.COMPONENT.value,
            required="False",
        ),

        # GrammarElementFieldString(
        #     parent=None,
        #     title=RequirementFieldName.LEVEL,
        #     required="False",
        # ),
        # GrammarElementFieldString(
        #     parent=None,
        #     title=RequirementFieldName.STATUS,
        #     required="False",
        # ),
        # GrammarElementFieldString(
        #     parent=None, title=RequirementFieldName.TAGS, required="False"
        # ),
        GrammarElementFieldReference(
            parent=None,
            title=RequirementFieldName.REFS,
            types=[
                GrammarReferenceType.PARENT_REQ_REFERENCE,
            ],
            required="False",
        ),
        GrammarElementFieldString(
            parent=None,
            title=RequirementFieldName.TITLE,
            required="False",
        ),
        GrammarElementFieldString(
            parent=None,
            title=RequirementFieldName.STATEMENT,
            required="False",
        ),
        GrammarElementFieldString(
            parent=None,
            title=FIELDS.USER_STORY.value,
            required="False",
        ),
        GrammarElementFieldString(
            parent=None,
            title=FIELDS.DISCUSSION_DATE.value,
            required="False",
        ),
        GrammarElementFieldString(
            parent=None,
            title=FIELDS.REVIEW_COMMENT.value,
            required="False",
        ),
    ]
    requirement_element = GrammarElement(
        parent=None, tag="REQUIREMENT", fields=fields
    )
    elements: List[GrammarElement] = [requirement_element]

    document_grammar = DocumentGrammar(parent=document, elements=elements)
    document.grammar = document_grammar
    document_reference = DocumentReference()
    document_reference.set_document(document)

    for requirement_dict_ in requirements_dicts:
        requirement = Requirement(document, "REQUIREMENT", fields=[])
        requirement.ng_level = 1
        requirement.ng_document_reference = document_reference

        if requirement_dict_[FIELDS.UID.value] is not None:
            requirement.set_field_value(
                field_name=FIELDS.UID.value,
                form_field_index=0,
                value=requirement_dict_[FIELDS.UID.value],
            )
        if requirement_dict_[FIELDS.TYPE.value] is not None:
            requirement.set_field_value(
                field_name=FIELDS.TYPE.value,
                form_field_index=0,
                value=requirement_dict_[FIELDS.TYPE.value],
            )
        if requirement_dict_[FIELDS.COMPONENT.value] is not None:
            requirement.set_field_value(
                field_name=FIELDS.COMPONENT.value,
                form_field_index=0,
                value=requirement_dict_[FIELDS.COMPONENT.value].strip(),
            )
        if requirement_dict_[FIELDS.TITLE.value] is not None:
            requirement.set_field_value(
                field_name=FIELDS.TITLE.value,
                form_field_index=0,
                value=requirement_dict_[FIELDS.TITLE.value],
            )
        requirement.set_field_value(
            field_name=FIELDS.STATEMENT.value,
            form_field_index=0,
            value=requirement_dict_[FIELDS.STATEMENT.value],
        )
        if requirement_dict_[FIELDS.USER_STORY.value] is not None:
            requirement.set_field_value(
                field_name=FIELDS.USER_STORY.value,
                form_field_index=0,
                value=requirement_dict_[FIELDS.USER_STORY.value],
            )
        if requirement_dict_[FIELDS.DISCUSSION_DATE.value] is not None:
            requirement.set_field_value(
                field_name=FIELDS.DISCUSSION_DATE.value,
                form_field_index=0,
                # Some cells have strings like 20221122. They get parsed as float().
                # Converting them to back to str().
                value=str(requirement_dict_[FIELDS.DISCUSSION_DATE.value]),
            )
        if requirement_dict_[FIELDS.REVIEW_COMMENT.value] is not None:
            requirement.set_field_value(
                field_name=FIELDS.REVIEW_COMMENT.value,
                form_field_index=0,
                value=requirement_dict_[FIELDS.REVIEW_COMMENT.value],
            )

        if (
            requirement_dict_[FIELDS.PARENT_UID.value] is not None
            and
            "--" not in requirement_dict_[FIELDS.PARENT_UID.value]
        ):
            references = [
                ParentReqReference(
                    parent=requirement, ref_uid=requirement_dict_[FIELDS.PARENT_UID.value]
                )
            ]
            requirement.ordered_fields_lookup[RequirementFieldName.REFS] = [
                RequirementField(
                    parent=requirement,
                    field_name=RequirementFieldName.REFS,
                    field_value=None,
                    field_value_multiline=None,
                    field_value_references=references,
                )
            ]
            requirement.references = references

        document.section_contents.append(requirement)

    print(f"Requirements total: {len(document.section_contents)}")
    sdoc_content = SDWriter().write(document)

    Path("docs").mkdir(parents=True, exist_ok=True)
    path_to_output_sdoc = os.path.join("docs", "zephyr.sdoc")
    with open(path_to_output_sdoc, "w") as output_sdoc_file:
        output_sdoc_file.write(sdoc_content)


main()
